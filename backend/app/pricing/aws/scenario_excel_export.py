"""Run AWS scenario pricing and export a single-sheet Excel workbook."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.clients.ai_client import AIClientFactory, BaseAIClient
from app.models import Project
from app.pricing.aws.benchmark import build_usage_service, normalize_aws_components
from app.pricing.aws.cost_calculator import AwsCostCalculator
from app.pricing.aws.project_costing import AwsProjectCostingPipeline
from app.pricing.aws.project_pricing import calculate_project_aws_sku_quantities
from app.pricing.aws.registry import get_aws_pricing_model
from app.pricing.aws.scenarios import AWS_TWENTY_SCENARIOS, twenty_scenario_components
from app.pricing.aws.usage_inference import AwsUsageInferenceEngine
from app.pricing.azure.benchmark import USER_COUNT_TO_PROJECT_LABEL
from app.pricing.azure.scenarios import BenchmarkScenario
from app.pricing.schemas import (
    ComponentCostResult,
    ComponentPricingInput,
    ComponentSkuResult,
    SkuCostLine,
    UsageAssumption,
)
from app.pricing.usage.service import UsageAssumptionsService

InferenceMode = Literal["llm", "heuristic"]
CatalogMode = Literal["firestore"]

EXCEL_HEADERS: tuple[str, ...] = (
    "product name",
    "component name",
    "users",
    "price",
    "llm usage assumptions",
    "how calculated",
    "calculated_skus",
)


@dataclass(frozen=True)
class ScenarioComponentRow:
    product_name: str
    component_name: str
    component_id: str
    aws_service: str
    users: int
    price_usd: float
    llm_usage_assumptions: str
    how_calculated: str
    skus: str
    inference_source: str
    pricing_status: str


@dataclass(frozen=True)
class ScenarioExportResult:
    scenario_id: str
    product_name: str
    users: int
    total_usd: float
    inference_source: str
    rows: tuple[ScenarioComponentRow, ...]
    warnings: tuple[str, ...]


def build_aws_cost_calculator(*, catalog: CatalogMode = "firestore") -> AwsCostCalculator:
    """Build AWS cost calculator from Firestore catalog (single source of truth)."""
    del catalog
    from app.pricing.catalog_factory import build_aws_cost_calculator as _build

    return _build()


def _format_assumptions(
    behavioral: list[UsageAssumption],
    resolved: list[UsageAssumption],
) -> str:
    parts: list[str] = []
    if behavioral:
        parts.append("Behavioral:")
        for item in behavioral:
            unit = f" {item.unit}" if item.unit else ""
            parts.append(f"{item.key}={item.value}{unit} ({item.confidence.value})")
    if resolved:
        if parts:
            parts.append("")
        parts.append("Resolved:")
        for item in resolved:
            unit = f" {item.unit}" if item.unit else ""
            parts.append(f"{item.key}={item.value}{unit} ({item.source.value})")
    return "\n".join(parts)


def _format_how_calculated(
    sku_result: ComponentSkuResult | None,
    line_items: list[SkuCostLine],
) -> str:
    parts: list[str] = []
    if sku_result is not None:
        for quantity in sku_result.quantities:
            if quantity.quantity <= 0:
                continue
            model = get_aws_pricing_model(sku_result.service)
            derivation = quantity.formula
            if model is not None:
                sku_def = next(
                    (item for item in model.pricing_model.calculated_skus if item.key == quantity.sku_key),
                    None,
                )
                if sku_def is not None and sku_def.derivation:
                    derivation = sku_def.derivation
            inputs = ", ".join(f"{key}={value}" for key, value in quantity.input_values_used.items())
            parts.append(
                f"{quantity.sku_key}: {derivation} "
                f"[qty={quantity.quantity:g} {quantity.unit}; inputs: {inputs}]"
            )
    for line in line_items:
        if line.formula_note:
            parts.append(f"{line.sku_key} billing: {line.formula_note}")
    return "\n".join(parts)


def _format_skus(line_items: list[SkuCostLine]) -> str:
    if not line_items:
        return ""
    parts: list[str] = []
    for line in line_items:
        parts.append(
            f"{line.sku_key}/{line.catalog_role}: "
            f"{line.billable_units:g} {line.usage_unit} "
            f"@ ${line.unit_price_usd:g} = ${line.monthly_cost_usd:,.2f}"
        )
    return "\n".join(parts)


class AwsScenarioExcelExporter:
    """Run twenty AWS scenarios through usage inference and catalog pricing."""

    def __init__(
        self,
        pipeline: AwsProjectCostingPipeline,
        usage_service: UsageAssumptionsService,
        *,
        inference_mode: InferenceMode = "llm",
        users: int = 1_000,
        scenarios: tuple[BenchmarkScenario, ...] = AWS_TWENTY_SCENARIOS,
    ) -> None:
        self._pipeline = pipeline
        self._usage_service = usage_service
        self._inference_mode = inference_mode
        self._users = users
        self._scenarios = scenarios

    def run_all(self) -> list[ScenarioExportResult]:
        results: list[ScenarioExportResult] = []
        for index, scenario in enumerate(self._scenarios, start=1):
            print(f"[{index}/{len(self._scenarios)}] {scenario.name} ({self._users:,} users)...")
            results.append(self._run_scenario(scenario))
        return results

    def _run_scenario(self, scenario: BenchmarkScenario) -> ScenarioExportResult:
        components = normalize_aws_components(twenty_scenario_components(scenario))
        component_names = {spec.key: spec.name for spec in scenario.components}

        project = Project(
            name=scenario.name,
            description=scenario.product_description,
            expected_users=USER_COUNT_TO_PROJECT_LABEL[self._users],
            stage=scenario.stage,
        )
        inference = self._usage_service.infer(
            project,
            components,
            provider="aws",
            feature_flags=scenario.feature_flags,
            inference_mode=self._inference_mode,
        )
        pricing_inputs = list(inference.components)
        quantity_result = calculate_project_aws_sku_quantities(pricing_inputs)
        cost_result = self._pipeline.calculate(
            project,
            components,
            pricing_inputs=pricing_inputs,
            feature_flags=scenario.feature_flags,
        )

        sku_by_id = {item.component_id: item for item in quantity_result.components}
        cost_by_id = {item.component_id: item for item in cost_result.components}
        input_by_id = {item.component_id: item for item in pricing_inputs}

        rows: list[ScenarioComponentRow] = []
        for component in components:
            component_id = component.key
            cost_component = cost_by_id.get(component_id)
            pricing_input = input_by_id.get(component_id)
            sku_component = sku_by_id.get(component_id)
            rows.append(
                self._build_row(
                    scenario=scenario,
                    component_id=component_id,
                    component_name=component_names.get(component_id, component.name),
                    aws_service=str(component.cloud.get("aws", "")),
                    cost_component=cost_component,
                    pricing_input=pricing_input,
                    sku_component=sku_component,
                    inference_source=inference.inference_source,
                )
            )

        return ScenarioExportResult(
            scenario_id=scenario.scenario_id,
            product_name=scenario.name,
            users=self._users,
            total_usd=round(cost_result.total_usd, 2),
            inference_source=inference.inference_source,
            rows=tuple(rows),
            warnings=tuple(cost_result.warnings),
        )

    def _build_row(
        self,
        *,
        scenario: BenchmarkScenario,
        component_id: str,
        component_name: str,
        aws_service: str,
        cost_component: ComponentCostResult | None,
        pricing_input: ComponentPricingInput | None,
        sku_component: ComponentSkuResult | None,
        inference_source: str,
    ) -> ScenarioComponentRow:
        line_items = cost_component.line_items if cost_component else []
        behavioral = pricing_input.behavioral_assumptions if pricing_input else []
        resolved = pricing_input.resolved if pricing_input else []
        if cost_component is not None:
            behavioral = cost_component.behavioral_assumptions or behavioral
            resolved = cost_component.resolved_assumptions or resolved

        return ScenarioComponentRow(
            product_name=scenario.name,
            component_name=component_name,
            component_id=component_id,
            aws_service=aws_service,
            users=self._users,
            price_usd=round(cost_component.subtotal_usd, 2) if cost_component else 0.0,
            llm_usage_assumptions=_format_assumptions(behavioral, resolved),
            how_calculated=_format_how_calculated(sku_component, line_items),
            skus=_format_skus(line_items),
            inference_source=inference_source,
            pricing_status=cost_component.pricing_status if cost_component else "unknown",
        )


def write_scenario_excel(
    results: list[ScenarioExportResult],
    output_path: Path,
    *,
    inference_mode: InferenceMode,
    catalog: CatalogMode,
) -> Path:
    """Write all scenario component tables to a single Excel sheet."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "AWS Scenario Pricing"

    meta_font = Font(bold=True)
    sheet.append(["AWS Scenario Pricing Export"])
    sheet.append([f"Generated: {datetime.now(tz=UTC).isoformat()}"])
    sheet.append([f"Inference mode: {inference_mode}"])
    sheet.append([f"Catalog: {catalog}"])
    sheet.append([f"Scenarios: {len(results)}"])
    sheet.append([])

    for index, scenario_result in enumerate(results):
        if index > 0:
            sheet.append([])

        sheet.append([f"Scenario: {scenario_result.product_name}"])
        sheet.append(
            [
                f"Users: {scenario_result.users:,}",
                f"Total: ${scenario_result.total_usd:,.2f}",
                f"Inference: {scenario_result.inference_source}",
            ]
        )
        header_row = sheet.max_row + 1
        sheet.append(list(EXCEL_HEADERS))
        for cell in sheet[header_row]:
            cell.font = meta_font

        for row in scenario_result.rows:
            sheet.append(
                [
                    row.product_name,
                    f"{row.component_name} ({row.aws_service})",
                    row.users,
                    row.price_usd,
                    row.llm_usage_assumptions,
                    row.how_calculated,
                    row.skus,
                ]
            )

        total_row = sheet.max_row + 1
        sheet.append(
            [
                scenario_result.product_name,
                "TOTAL",
                scenario_result.users,
                scenario_result.total_usd,
                "",
                "",
                "",
            ]
        )
        for cell in sheet[total_row]:
            cell.font = meta_font

    _autosize_columns(sheet)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _autosize_columns(sheet) -> None:
    for column_index, header in enumerate(EXCEL_HEADERS, start=1):
        letter = get_column_letter(column_index)
        max_length = len(header)
        for row in sheet.iter_rows(min_col=column_index, max_col=column_index, values_only=True):
            if row[0] is None:
                continue
            max_length = max(max_length, min(len(str(row[0])), 80))
        sheet.column_dimensions[letter].width = max_length + 2


def build_exporter(
    *,
    inference_mode: InferenceMode = "llm",
    catalog: CatalogMode = "firestore",
    users: int = 1_000,
    ai_client: BaseAIClient | None = None,
) -> AwsScenarioExcelExporter:
    calculator = build_aws_cost_calculator(catalog=catalog)
    pipeline = AwsProjectCostingPipeline(AwsUsageInferenceEngine(), calculator)
    if ai_client is None and inference_mode == "llm":
        ai_client = AIClientFactory.create()
    usage_service = build_usage_service(ai_client=ai_client, inference_mode=inference_mode)
    return AwsScenarioExcelExporter(
        pipeline,
        usage_service,
        inference_mode=inference_mode,
        users=users,
    )
