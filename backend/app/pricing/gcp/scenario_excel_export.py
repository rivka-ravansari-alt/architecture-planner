"""Run GCP scenario pricing and export a multi-sheet Excel benchmark workbook."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.clients.ai_client import AIClientFactory, BaseAIClient
from app.models import Project
from app.pricing.azure.benchmark import BENCHMARK_USER_COUNTS, USER_COUNT_TO_PROJECT_LABEL
from app.pricing.azure.scenarios import (
    AI_CHAT,
    AI_DOCUMENT_OCR,
    ECOMMERCE,
    SELF_ESTEEM_HABIT,
    SIMPLE_CRUD_SAAS,
    BenchmarkScenario,
)
from app.pricing.gcp.benchmark import build_usage_service, normalize_gcp_components
from app.pricing.gcp.cost_calculator import GcpCostCalculator
from app.pricing.gcp.scenarios import gcp_scenario_components
from app.pricing.gcp.verification import (
    GcpPricingVerificationReport,
    GcpPricingVerificationRunner,
    build_validation_checks,
    pricing_completeness,
)

InferenceMode = Literal["llm", "heuristic"]
CatalogMode = Literal["firestore"]

GCP_FIVE_SCENARIOS: tuple[BenchmarkScenario, ...] = (
    SIMPLE_CRUD_SAAS,
    SELF_ESTEEM_HABIT,
    ECOMMERCE,
    AI_CHAT,
    AI_DOCUMENT_OCR,
)

AWS_STYLE_HEADERS: tuple[str, ...] = (
    "product name",
    "component name",
    "users",
    "price",
    "llm usage assumptions",
    "how calculated",
    "calculated_skus",
)


@dataclass(frozen=True)
class ScenarioRunExport:
    scenario: BenchmarkScenario
    users: int
    report: GcpPricingVerificationReport


def build_gcp_cost_calculator(*, catalog: CatalogMode = "firestore") -> GcpCostCalculator:
    """Build GCP cost calculator from Firestore catalog (single source of truth)."""
    del catalog
    from app.pricing.catalog_factory import build_gcp_cost_calculator as _build

    return _build()


def _format_assumptions(report: GcpPricingVerificationReport, component_id: str) -> str:
    component = next(
        (item for item in report.components if item.architecture.component_id == component_id),
        None,
    )
    if component is None:
        return ""
    parts: list[str] = []
    if component.behavioral_assumptions:
        parts.append("Per-user behavior:")
        for item in component.behavioral_assumptions:
            unit = f" {item.unit}" if item.unit else ""
            parts.append(
                f"{item.key}={item.value}{unit} ({item.confidence.value}) — {item.reasoning}"
            )
    if component.resolved_assumptions:
        if parts:
            parts.append("")
        parts.append("Resolved / infrastructure:")
        for item in component.resolved_assumptions:
            unit = f" {item.unit}" if item.unit else ""
            parts.append(
                f"{item.key}={item.value}{unit} ({item.source.value}, {item.confidence.value})"
            )
    return "\n".join(parts)


def _format_how_calculated(component) -> str:
    parts: list[str] = []
    for line in component.sku_lines:
        if line.formula:
            inputs = ", ".join(f"{k}={v}" for k, v in line.input_values_used.items())
            parts.append(
                f"{line.sku_key}: {line.formula} "
                f"[raw={line.raw_quantity or 0:g} {line.unit}; billable={line.billable_quantity:g}; inputs: {inputs}]"
            )
        if line.free_tier_deducted > 0:
            parts.append(f"{line.sku_key}: free_tier_deducted={line.free_tier_deducted:g}")
    return "\n".join(parts)


def _format_skus(component) -> str:
    parts: list[str] = []
    for line in component.sku_lines:
        if line.sku_cost_usd is None:
            continue
        parts.append(
            f"{line.sku_key}/{line.catalog_role}: "
            f"{line.billable_units or 0:g} {line.catalog_usage_unit or ''} "
            f"@ ${line.unit_price_usd or 0:g} = ${line.sku_cost_usd:,.2f}"
        )
    return "\n".join(parts)


class GcpScenarioExcelExporter:
    """Run five GCP benchmark scenarios across user counts and export Excel."""

    def __init__(
        self,
        runner: GcpPricingVerificationRunner,
        usage_service,
        *,
        inference_mode: InferenceMode = "heuristic",
        user_counts: tuple[int, ...] = BENCHMARK_USER_COUNTS,
        scenarios: tuple[BenchmarkScenario, ...] = GCP_FIVE_SCENARIOS,
    ) -> None:
        self._runner = runner
        self._usage_service = usage_service
        self._inference_mode = inference_mode
        self._user_counts = user_counts
        self._scenarios = scenarios

    def run_all(self) -> list[ScenarioRunExport]:
        results: list[ScenarioRunExport] = []
        total_runs = len(self._scenarios) * len(self._user_counts)
        run_index = 0
        for scenario in self._scenarios:
            components = normalize_gcp_components(gcp_scenario_components(scenario))
            for users in self._user_counts:
                run_index += 1
                print(
                    f"[{run_index}/{total_runs}] {scenario.name} ({users:,} users)...",
                    flush=True,
                )
                project = Project(
                    name=scenario.name,
                    description=scenario.product_description,
                    expected_users=USER_COUNT_TO_PROJECT_LABEL[users],
                    stage=scenario.stage,
                )
                inference = self._usage_service.infer(
                    project,
                    components,
                    provider="gcp",
                    feature_flags=scenario.feature_flags,
                    inference_mode=self._inference_mode,
                )
                report = self._runner.run(
                    project,
                    components,
                    feature_flags=scenario.feature_flags,
                    pricing_inputs=list(inference.components),
                    inference_source=inference.inference_source,
                )
                results.append(
                    ScenarioRunExport(
                        scenario=scenario,
                        users=users,
                        report=report,
                    )
                )
        return results


def write_gcp_scenario_excel(
    runs: list[ScenarioRunExport],
    output_path: Path,
    *,
    inference_mode: InferenceMode,
    catalog: CatalogMode,
) -> Path:
    workbook = Workbook()
    bold = Font(bold=True)

    summary = workbook.active
    summary.title = "Summary"
    summary.append(["GCP Scenario Pricing Benchmark"])
    summary.append([f"Generated: {datetime.now(tz=UTC).isoformat()}"])
    summary.append([f"Inference mode: {inference_mode}"])
    summary.append([f"Catalog: {catalog}"])
    summary.append([f"Scenarios: {len(GCP_FIVE_SCENARIOS)}"])
    summary.append([f"User counts: {', '.join(str(u) for u in BENCHMARK_USER_COUNTS)}"])
    summary.append([])

    user_counts = BENCHMARK_USER_COUNTS
    header = ["Scenario"] + [str(u) for u in user_counts]
    summary.append(header)
    for cell in summary[summary.max_row]:
        cell.font = bold

    totals_by_scenario: dict[str, dict[int, float]] = {}
    for run in runs:
        totals_by_scenario.setdefault(run.scenario.name, {})[run.users] = run.report.total_usd

    for scenario in GCP_FIVE_SCENARIOS:
        row = [scenario.name]
        row.extend(f"${totals_by_scenario.get(scenario.name, {}).get(u, 0.0):,.2f}" for u in user_counts)
        summary.append(row)

    summary.append([])
    summary.append(["Monthly total USD by scenario and user count"])
    _autosize_sheet(summary)

    arch = workbook.create_sheet("Architecture Mapping")
    arch.append(
        ["scenario", "users", "component_id", "component_name", "component_type", "gcp_service"]
    )
    for cell in arch[1]:
        cell.font = bold

    usage = workbook.create_sheet("Usage Assumptions")
    usage.append(
        [
            "scenario",
            "users",
            "component",
            "assumption_type",
            "key",
            "value",
            "unit",
            "confidence",
            "reasoning",
        ]
    )
    for cell in usage[1]:
        cell.font = bold

    sku_qty = workbook.create_sheet("SKU Quantities")
    sku_qty.append(
        [
            "scenario",
            "users",
            "component",
            "gcp_service",
            "sku_key",
            "raw_quantity",
            "free_tier_deducted",
            "billable_quantity",
            "unit",
            "formula",
        ]
    )
    for cell in sku_qty[1]:
        cell.font = bold

    catalog_sheet = workbook.create_sheet("Catalog Pricing")
    catalog_sheet.append(
        [
            "scenario",
            "users",
            "component",
            "gcp_service",
            "sku_key",
            "catalog_role",
            "billable_units",
            "usage_unit",
            "unit_price_usd",
            "monthly_cost_usd",
            "missing_price_reason",
        ]
    )
    for cell in catalog_sheet[1]:
        cell.font = bold

    breakdown = workbook.create_sheet("Component Breakdown")
    breakdown.append(
        [
            "scenario",
            "users",
            "component_id",
            "component_name",
            "gcp_service",
            "subtotal_usd",
            "pricing_status",
            "ready",
        ]
    )
    for cell in breakdown[1]:
        cell.font = bold

    validation = workbook.create_sheet("Validation")
    validation.append(
        [
            "scenario",
            "users",
            "total_usd",
            "pricing_completeness",
            "missing_prices",
            "warnings_count",
            "rule_id",
            "status",
            "message",
        ]
    )
    for cell in validation[1]:
        cell.font = bold

    pricing = workbook.create_sheet("GCP Scenario Pricing")
    pricing.append(["GCP Scenario Pricing Export"])
    pricing.append([f"Generated: {datetime.now(tz=UTC).isoformat()}"])
    pricing.append([f"Inference mode: {inference_mode}"])
    pricing.append([f"Catalog: {catalog}"])
    pricing.append([])

    for run in runs:
        scenario = run.scenario
        report = run.report
        spec_by_id = {spec.key: spec for spec in scenario.components}

        for spec in scenario.components:
            arch.append(
                [
                    scenario.name,
                    run.users,
                    spec.key,
                    spec.name,
                    spec.component_type,
                    spec.gcp_service,
                ]
            )

        for component in report.components:
            comp_label = f"{component.architecture.component_name} ({component.gcp_service})"
            for assumption in component.behavioral_assumptions:
                usage.append(
                    [
                        scenario.name,
                        run.users,
                        comp_label,
                        "behavioral",
                        assumption.key,
                        assumption.value,
                        assumption.unit or "",
                        assumption.confidence.value,
                        assumption.reasoning,
                    ]
                )
            for assumption in component.resolved_assumptions:
                usage.append(
                    [
                        scenario.name,
                        run.users,
                        comp_label,
                        "resolved",
                        assumption.key,
                        assumption.value,
                        assumption.unit or "",
                        assumption.confidence.value,
                        assumption.reasoning,
                    ]
                )

            for line in component.sku_lines:
                sku_qty.append(
                    [
                        scenario.name,
                        run.users,
                        comp_label,
                        component.gcp_service,
                        line.sku_key,
                        line.raw_quantity,
                        line.free_tier_deducted,
                        line.billable_quantity,
                        line.unit,
                        line.formula,
                    ]
                )
                catalog_sheet.append(
                    [
                        scenario.name,
                        run.users,
                        comp_label,
                        component.gcp_service,
                        line.sku_key,
                        line.catalog_role or "",
                        line.billable_units,
                        line.catalog_usage_unit or "",
                        line.unit_price_usd,
                        line.sku_cost_usd,
                        line.missing_price_reason or "",
                    ]
                )

            breakdown.append(
                [
                    scenario.name,
                    run.users,
                    component.architecture.component_id,
                    component.architecture.component_name,
                    component.gcp_service,
                    round(component.subtotal_usd, 2),
                    component.pricing_status,
                    component.ready,
                ]
            )

        completeness = pricing_completeness(report)
        checks = build_validation_checks(report)
        for check in checks:
            validation.append(
                [
                    scenario.name,
                    run.users,
                    report.total_usd,
                    completeness,
                    len(report.missing_prices),
                    len(report.warnings),
                    check["rule_id"],
                    check["status"],
                    check["message"],
                ]
            )

        pricing.append([])
        pricing.append([f"Scenario: {scenario.name}"])
        pricing.append(
            [
                f"Users: {run.users:,}",
                f"Total: ${report.total_usd:,.2f}",
                f"Inference: {report.inference_source}",
                f"Completeness: {completeness}",
            ]
        )
        header_row = pricing.max_row + 1
        pricing.append(list(AWS_STYLE_HEADERS))
        for cell in pricing[header_row]:
            cell.font = bold

        for component in report.components:
            spec = spec_by_id.get(component.architecture.component_id)
            component_name = component.architecture.component_name
            if spec is not None:
                component_name = spec.name
            pricing.append(
                [
                    scenario.name,
                    f"{component_name} ({component.gcp_service})",
                    run.users,
                    round(component.subtotal_usd, 2),
                    _format_assumptions(report, component.architecture.component_id),
                    _format_how_calculated(component),
                    _format_skus(component),
                ]
            )

        total_row = pricing.max_row + 1
        pricing.append(
            [
                scenario.name,
                "TOTAL",
                run.users,
                report.total_usd,
                "",
                "",
                "",
            ]
        )
        for cell in pricing[total_row]:
            cell.font = bold

    for sheet in (arch, usage, sku_qty, catalog_sheet, breakdown, validation):
        _autosize_sheet(sheet)
    _autosize_sheet(pricing, max_width=80)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _autosize_sheet(sheet, *, max_width: int = 50) -> None:
    for column_index in range(1, sheet.max_column + 1):
        letter = get_column_letter(column_index)
        max_length = 10
        for row in sheet.iter_rows(min_col=column_index, max_col=column_index, values_only=True):
            if row[0] is None:
                continue
            max_length = max(max_length, min(len(str(row[0])), max_width))
        sheet.column_dimensions[letter].width = max_length + 2


def build_exporter(
    *,
    inference_mode: InferenceMode = "heuristic",
    catalog: CatalogMode = "firestore",
    ai_client: BaseAIClient | None = None,
) -> GcpScenarioExcelExporter:
    calculator = build_gcp_cost_calculator(catalog=catalog)
    runner = GcpPricingVerificationRunner(calculator, inference_mode=inference_mode)
    if ai_client is None and inference_mode == "llm":
        ai_client = AIClientFactory.create()
    usage_service = build_usage_service(ai_client=ai_client, inference_mode=inference_mode)
    return GcpScenarioExcelExporter(runner, usage_service, inference_mode=inference_mode)
